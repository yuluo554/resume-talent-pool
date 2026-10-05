"""生成器单测（M1 DoD③）：组数/同名组/缺失率/seed 复现 + 三格式抽取 roundtrip。

roundtrip 把每份生成文件的文本抽取出来（pdfplumber / python-docx / 解码），
归一化（全角→半角、去空白、小写）后断言 truth 声明的每个值都出现在原文中——
同时充当"生成 PDF/docx 字体可抽取"的自动化门（DoD④ 的程序化部分）。
日期按渲染变体匹配（YYYY-MM / YYYY年M月 等价）。

依赖纪律：本模块真实 import reportlab / python-docx / pdfplumber（gen+parse extras），
CI 必须安装 `. [dev,parse,gen]` 并对照收集数，禁止 importorskip 整模块跳过。
"""

import json
import re
from pathlib import Path

import pytest

from resume_talent_pool.cli import main as cli_main
from resume_talent_pool.evaluation.generator import (
    DEFAULT_SEED,
    FIELD_KEYS,
    DEGREE_ENUM,
    generate,
    load_specs,
)

_DATE_RE = re.compile(r"\d{4}-\d{2}")
_FW_TRANS = str.maketrans("０１２３４５６７８９－", "0123456789-")


def _normalize(text: str) -> str:
    return re.sub(r"\s+", "", text).translate(_FW_TRANS).lower()


def _date_ok(text: str, iso: str) -> bool:
    y, m = iso[:4], int(iso[5:7])
    variants = [f"{y}-{m:02d}", f"{y}年{m}月", f"{y}-{m}", f"{y}.{m:02d}", f"{y}年{m:02d}月"]
    return any(v in text for v in variants)


def _value_ok(text: str, value) -> bool:
    if isinstance(value, str) and _DATE_RE.fullmatch(value):
        return _date_ok(text, value)
    return _normalize(str(value)) in text


def _extract_text(path: Path, fmt: str, encoding: str) -> str:
    if fmt == "txt":
        return path.read_bytes().decode(encoding)
    if fmt == "pdf":
        import pdfplumber
        with pdfplumber.open(str(path)) as pdf:
            return "\n".join(page.extract_text() or "" for page in pdf.pages)
    import docx
    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append("".join(cell.text for cell in row.cells))
    return "\n".join(parts)


@pytest.fixture(scope="session")
def run100(tmp_path_factory):
    out = tmp_path_factory.mktemp("gen100")
    summary = generate(str(out), seed=DEFAULT_SEED, persons=100)
    truth = json.loads(Path(summary["truth_path"]).read_text(encoding="utf-8"))
    return summary, truth, Path(summary["resumes_dir"])


@pytest.fixture(scope="session")
def run40(tmp_path_factory):
    out = tmp_path_factory.mktemp("gen40")
    summary = generate(str(out), seed=DEFAULT_SEED, persons=40)
    truth = json.loads(Path(summary["truth_path"]).read_text(encoding="utf-8"))
    return summary, truth, Path(summary["resumes_dir"])


# --------------------------------------------------------------------------- #
# 结构与组数
# --------------------------------------------------------------------------- #

def test_full_batch_shape(run100):
    summary, truth, resumes_dir = run100
    assert summary["persons"] == 100
    assert 160 <= summary["files"] <= 200  # ~60% 人 2-3 版本 → 约 180
    truth_files = [r["file"] for p in truth["persons"] for r in p["resumes"]]
    assert len(truth_files) == summary["files"] == len(set(truth_files))
    on_disk = {f.name for f in resumes_dir.iterdir()}
    assert on_disk == set(truth_files)
    n = summary["files"]
    assert 0.30 <= summary["by_format"]["pdf"] / n <= 0.50
    assert 0.30 <= summary["by_format"]["docx"] / n <= 0.50
    assert 0.12 <= summary["by_format"]["txt"] / n <= 0.30


def test_homonym_groups(run100):
    _, truth, _ = run100
    groups = truth["relations"]["homonym_groups"]
    assert len(groups) == 5
    specs = load_specs()
    pool = {(g["name_a"], g["name_b"]) for g in specs["names"]["homonym_groups"]}
    persons = {p["person_id"]: p for p in truth["persons"]}
    for pair in groups:
        assert len(pair) == 2
        a, b = persons[pair[0]], persons[pair[1]]
        names = (a["name"], b["name"])
        assert names in pool or names[::-1] in pool  # 同名/同音组与参数池对得上
        r1, r2 = a["resumes"][0], b["resumes"][0]
        # phone/email/经历完全独立（防误合并的显式反例）
        f1, f2 = {r["file"]: r["truth_fields"] for r in a["resumes"]}, \
                 {r["file"]: r["truth_fields"] for r in b["resumes"]}
        phones1 = {f["phone"] for f in f1.values() if "phone" in f}
        phones2 = {f["phone"] for f in f2.values() if "phone" in f}
        assert phones1.isdisjoint(phones2)
        emails1 = {f["email"] for f in f1.values() if "email" in f}
        emails2 = {f["email"] for f in f2.values() if "email" in f}
        assert emails1.isdisjoint(emails2)
        companies1 = {e["company"] for r in a["resumes"] for e in r["truth_experiences"]}
        companies2 = {e["company"] for r in b["resumes"] for e in r["truth_experiences"]}
        assert companies1.isdisjoint(companies2)


def test_same_person_grouping(run100):
    summary, truth, _ = run100
    persons = truth["persons"]
    groups = [list(g) for g in truth["relations"]["same_person_groups"]]
    # 每人恰属一组，组并集覆盖全体（聚类真值完整性）
    members = [pid for g in groups for pid in g]
    assert sorted(members) == sorted(p["person_id"] for p in persons)
    multi = sum(1 for p in persons if len(p["resumes"]) > 1)
    assert 0.50 <= multi / len(persons) <= 0.70  # ~60% 人有 2+ 版本
    assert summary["multi_version_persons"] == multi
    for p in persons:  # 应届生版本封顶 2
        if p["template"] == "fresh":
            assert len(p["resumes"]) <= 2
        assert p["same_person_group"] == p["person_id"]


def test_version_evolution(run100):
    _, truth, _ = run100
    changed_phone = changed_email = total_multi = 0
    for p in truth["persons"]:
        if len(p["resumes"]) < 2:
            continue
        total_multi += 1
        resumes = p["resumes"]
        first, last = resumes[0], resumes[-1]
        base = {k for k in first["truth_fields"]}
        for prev, curr in zip(resumes, resumes[1:]):
            # 技能只增不减；经历（公司,职位,起）只增不减
            assert set(curr["truth_skills"]) >= set(prev["truth_skills"])
            keys_prev = {(e["company"], e["title"], e["start"]) for e in prev["truth_experiences"]}
            keys_curr = {(e["company"], e["title"], e["start"]) for e in curr["truth_experiences"]}
            assert keys_curr >= keys_prev
            # 学历不变
            for field in ("degree", "school", "major", "graduation_date"):
                if field in prev["truth_fields"] and field in curr["truth_fields"]:
                    assert prev["truth_fields"][field] == curr["truth_fields"][field]
        # 经历时间线连续：进行中段在其后版本获得确定 end
        for prev, curr in zip(resumes, resumes[1:]):
            for e_prev in prev["truth_experiences"]:
                if e_prev["end"] is None:
                    match = [e for e in curr["truth_experiences"]
                             if (e["company"], e["title"], e["start"]) ==
                             (e_prev["company"], e_prev["title"], e_prev["start"])]
                    if match:
                        assert match[0]["end"] is not None
        # 换号/换邮箱率（30% 口径的观测带宽）
        f1, f2 = first["truth_fields"], last["truth_fields"]
        if "phone" in f1 and "phone" in f2 and f1["phone"] != f2["phone"]:
            changed_phone += 1
        if "email" in f1 and "email" in f2 and f1["email"] != f2["email"]:
            changed_email += 1
    assert total_multi > 0
    assert 0.12 <= changed_phone / total_multi <= 0.48
    assert 0.12 <= changed_email / total_multi <= 0.48


def test_missing_rates(run100):
    _, truth, _ = run100
    resumes = [r for p in truth["persons"] for r in p["resumes"]]
    n = len(resumes)
    assert all("name" in r["truth_fields"] for r in resumes)  # name 恒在
    for field in ("phone", "email", "degree", "school", "major", "graduation_date",
                  "desired_position"):
        missing = sum(1 for r in resumes if field not in r["truth_fields"]) / n
        assert 0.05 <= missing <= 0.40, f"{field} 缺失率 {missing:.2f} 越界（口径 10-30%）"
    empty_certs = sum(1 for r in resumes if not r["truth_certificates"]) / n
    assert 0.03 <= empty_certs <= 0.22


def test_noise_coverage(run100):
    _, truth, _ = run100
    resumes = [r for p in truth["persons"] for r in p["resumes"]]
    n = len(resumes)
    noisy = sum(1 for r in resumes if r["noise"]) / n
    assert 0.35 <= noisy <= 0.68  # ~50% 文件带噪声
    kinds = {k for r in resumes for k in r["noise"]}
    for kind in ("pseudo_space", "fullwidth_digit", "separator_variant",
                 "multiline_experience", "crlf", "gbk_encoding_txt", "docx_table"):
        assert kind in kinds, f"全量批次未覆盖噪声类型 {kind}"
    for r in resumes:  # 格式专属噪声不越界
        assert "gbk_encoding_txt" not in r["noise"] or r["format"] == "txt"
        assert "docx_table" not in r["noise"] or r["format"] == "docx"
        if r["format"] == "txt":
            assert r["encoding"] in ("utf-8", "gbk")


def test_truth_schema(run100):
    _, truth, _ = run100
    assert truth["seed"] == DEFAULT_SEED
    for p in truth["persons"]:
        assert re.fullmatch(r"p\d{3}", p["person_id"])
        for r in p["resumes"]:
            assert set(r["truth_fields"]) <= set(FIELD_KEYS)
            assert r["file"].endswith("." + r["format"])
            for key, value in r["truth_fields"].items():
                assert value not in (None, "")
                if key == "degree":
                    assert value in DEGREE_ENUM
                if _DATE_RE.fullmatch(str(value)):
                    assert re.fullmatch(r"(19|20)\d{2}-(0[1-9]|1[0-2])", value)
            for e in r["truth_experiences"]:
                assert e["company"] and e["title"] and _DATE_RE.fullmatch(e["start"])
                assert e["end"] is None or _DATE_RE.fullmatch(e["end"])
                for m in (e["start"], e["end"] or ""):
                    if m:
                        assert re.fullmatch(r"(19|20)\d{2}-(0[1-9]|1[0-2])", m)


# --------------------------------------------------------------------------- #
# seed 复现（除文件时间戳外一致——本实现连时间戳都是固定的，逐字节相等）
# --------------------------------------------------------------------------- #

def test_seed_reproducible_bytes(tmp_path, run40):
    _, _, _ = run40  # 触发同 seed 会话级生成
    out_a, out_b = tmp_path / "a", tmp_path / "b"
    generate(str(out_a), seed=DEFAULT_SEED, persons=40)
    generate(str(out_b), seed=DEFAULT_SEED, persons=40)
    assert (out_a / "truth.json").read_bytes() == (out_b / "truth.json").read_bytes()
    for f in sorted((out_a / "resumes").iterdir()):
        assert f.read_bytes() == (out_b / "resumes" / f.name).read_bytes(), f.name


def test_seed_changes_output(tmp_path):
    out = tmp_path / "s"
    generate(str(out), seed=1, persons=10)
    truth1 = (out / "truth.json").read_bytes()
    generate(str(out), seed=2, persons=10)
    assert (out / "truth.json").read_bytes() != truth1


# --------------------------------------------------------------------------- #
# roundtrip：每份文件抽取文本 → truth 全值可复核（含字体可抽取门）
# --------------------------------------------------------------------------- #

def test_roundtrip_all_files_extractable(run100):
    _, truth, resumes_dir = run100
    failures = []
    for p in truth["persons"]:
        for r in p["resumes"]:
            text = _normalize(_extract_text(resumes_dir / r["file"], r["format"],
                                            r.get("encoding", "utf-8")))
            for key, value in r["truth_fields"].items():
                if not _value_ok(text, value):
                    failures.append((r["file"], key, value))
            for e in r["truth_experiences"]:
                for key in ("company", "title"):
                    if not _value_ok(text, e[key]):
                        failures.append((r["file"], "experience." + key, e[key]))
                if not _date_ok(text, e["start"]):
                    failures.append((r["file"], "experience.start", e["start"]))
                if e["end"] is None and "至今" not in text:
                    failures.append((r["file"], "experience.end", "至今"))
                elif e["end"] is not None and not _date_ok(text, e["end"]):
                    failures.append((r["file"], "experience.end", e["end"]))
            for item, kind in [(s, "skill") for s in r["truth_skills"]] + \
                              [(c, "certificate") for c in r["truth_certificates"]]:
                if not _value_ok(text, item):
                    failures.append((r["file"], kind, item))
    assert not failures, f"{len(failures)} 处真值不可从原文复核，示例: {failures[:10]}"


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #

def test_cli_bench_gen(tmp_path, capsys):
    out = tmp_path / "out"
    assert cli_main(["bench", "gen", "--seed", "7", "--persons", "3", "--out", str(out)]) == 0
    text = capsys.readouterr().out
    assert "生成完成" in text
    truth = json.loads((out / "truth.json").read_text(encoding="utf-8"))
    assert truth["seed"] == 7
    assert len(truth["persons"]) == 3
    assert (out / "resumes").is_dir()
